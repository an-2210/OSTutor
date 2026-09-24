"""
Document Ingestion Module for OSTutorLLM RAG Pipeline.

Discovers documents (.pdf, .pptx, .txt, .docx) in raw directory,
validates metadata against taxonomy, extracts text per page/slide,
and returns structured document page/slide records.
"""

import csv
import json
import logging
import os
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

# Third-party document extractors
try:
    import pypdf
except ImportError:
    pypdf = None

try:
    import pptx
except ImportError:
    pptx = None

try:
    import docx
except ImportError:
    docx = None

from config import METADATA_CSV_PATH, RAW_DATA_DIR, SUPPORTED_EXTENSIONS, TAXONOMY_JSON_PATH

logger = logging.getLogger(__name__)


def extract_pdf(path: Path) -> List[Tuple[int, str]]:
    """
    Extract text page by page from a PDF file using pypdf.

    Args:
        path: Path to the PDF file.

    Returns:
        List of (page_number_1_indexed, extracted_text) tuples.
    """
    if pypdf is None:
        raise ImportError("pypdf is required to extract text from PDF files. Install via pip install pypdf.")

    results: List[Tuple[int, str]] = []
    reader = pypdf.PdfReader(str(path))

    for idx, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        results.append((idx, text))

    return results


def extract_pptx(path: Path) -> List[Tuple[int, str]]:
    """
    Extract text slide by slide from a PPTX presentation using python-pptx.

    Args:
        path: Path to the PPTX file.

    Returns:
        List of (slide_number_1_indexed, extracted_text) tuples.
    """
    if pptx is None:
        raise ImportError("python-pptx is required to extract text from PPTX files. Install via pip install python-pptx.")

    results: List[Tuple[int, str]] = []
    prs = pptx.Presentation(str(path))

    for idx, slide in enumerate(prs.slides, start=1):
        slide_texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line:
                        slide_texts.append(line)
        text = "\n".join(slide_texts)
        results.append((idx, text))

    return results


def extract_txt(path: Path) -> List[Tuple[int, str]]:
    """
    Extract text from a plain text file.

    Args:
        path: Path to the TXT file.

    Returns:
        List containing a single (1, full_text) tuple.
    """
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        text = f.read()
    return [(1, text)]


def extract_docx(path: Path) -> List[Tuple[int, str]]:
    """
    Extract text from a Microsoft Word DOCX file using python-docx.

    Args:
        path: Path to the DOCX file.

    Returns:
        List containing a single (1, full_text) tuple.
    """
    if docx is None:
        raise ImportError("python-docx is required to extract text from DOCX files. Install via pip install python-docx.")

    doc = docx.Document(str(path))
    full_text = []
    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text.strip())
    return [(1, "\n".join(full_text))]


def extract_document(path: Path) -> List[Tuple[int, str]]:
    """
    Route document text extraction based on file extension.

    Args:
        path: Path to target document file.

    Returns:
        List of (page_or_slide_number, extracted_text) tuples.
    """
    ext = path.suffix.lower()
    if ext == ".pdf":
        return extract_pdf(path)
    elif ext == ".pptx":
        return extract_pptx(path)
    elif ext == ".txt":
        return extract_txt(path)
    elif ext == ".docx":
        return extract_docx(path)
    else:
        raise ValueError(f"Unsupported document format: {ext}")


def load_taxonomy(taxonomy_json_path: Path = TAXONOMY_JSON_PATH) -> Tuple[Set[int], Set[str]]:
    """
    Load valid units and topic names from data/taxonomy.json for validation.

    Returns:
        Tuple of (set of valid unit integers, set of valid topic names).
    """
    valid_units: Set[int] = set()
    valid_topics: Set[str] = set()

    if not taxonomy_json_path.exists():
        logger.warning(f"Taxonomy file missing at {taxonomy_json_path}. Skipping validation.")
        return valid_units, valid_topics

    with open(taxonomy_json_path, "r", encoding="utf-8") as f:
        units_data = json.load(f)

    for u in units_data:
        # Extract integer unit ID (e.g. "unit_1" -> 1)
        uid_str = u.get("unit_id", "")
        if uid_str.startswith("unit_"):
            try:
                valid_units.add(int(uid_str.split("_")[1]))
            except ValueError:
                pass
        for t in u.get("topics", []):
            if "name" in t:
                valid_topics.add(t["name"].lower())

    return valid_units, valid_topics


def load_and_validate_metadata(
    metadata_csv_path: Path = METADATA_CSV_PATH,
    taxonomy_json_path: Path = TAXONOMY_JSON_PATH,
) -> Dict[str, Dict[str, Any]]:
    """
    Load metadata catalog from metadata.csv and validate against taxonomy.

    Returns:
        Dictionary mapping filename and document_id to metadata dictionaries.
    """
    catalog: Dict[str, Dict[str, Any]] = {}
    valid_units, valid_topics = load_taxonomy(taxonomy_json_path)

    if not metadata_csv_path.exists():
        logger.warning(f"Metadata CSV missing at {metadata_csv_path}.")
        return catalog

    with open(metadata_csv_path, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            doc_id = row.get("document_id", "").strip()
            filename = row.get("filename", "").strip()

            # Convert unit to int
            unit_val = None
            try:
                unit_val = int(row.get("unit", ""))
            except ValueError:
                logger.warning(f"Metadata entry '{doc_id}': Invalid unit '{row.get('unit')}'.")

            # Validate unit
            if valid_units and unit_val not in valid_units:
                logger.warning(f"Metadata entry '{doc_id}': Unit {unit_val} not in syllabus taxonomy.")

            # Validate topic
            topic_val = row.get("topic", "").strip()
            if valid_topics and topic_val.lower() not in valid_topics:
                logger.warning(f"Metadata entry '{doc_id}': Topic '{topic_val}' not matched in syllabus taxonomy.")

            meta_record = {
                "document_id": doc_id or (filename.split(".")[0] if filename else "doc_unknown"),
                "filename": filename,
                "unit": unit_val or 1,
                "topic": topic_val or "General OS",
                "subtopic": row.get("subtopic", "").strip() or "General",
                "source_type": row.get("source_type", "").strip() or "faculty_material",
                "source": row.get("source", "").strip() or "Department Notes",
                "license_or_access": row.get("license_or_access", "").strip() or "Institutional Internal Use",
                "description": row.get("description", "").strip() or "",
            }

            if filename:
                catalog[filename.lower()] = meta_record
            if doc_id:
                catalog[doc_id.lower()] = meta_record

    return catalog


def discover_documents(root_dir: Path = RAW_DATA_DIR) -> List[Path]:
    """
    Scan root_dir recursively for supported document formats in deterministic order.

    Args:
        root_dir: Directory path to scan.

    Returns:
        Sorted list of Path objects for discovered valid documents.
    """
    discovered: List[Path] = []
    if not root_dir.exists():
        logger.warning(f"Raw document directory '{root_dir}' does not exist.")
        return discovered

    for path in sorted(root_dir.rglob("*")):
        if path.is_file():
            # Skip hidden files or temporary Office lock files (e.g. ~$slide.pptx)
            if path.name.startswith(".") or path.name.startswith("~$"):
                continue
            if path.suffix.lower() in SUPPORTED_EXTENSIONS:
                discovered.append(path)

    return sorted(discovered, key=lambda p: str(p).lower())


def ingest_document(
    file_path: Path, metadata_catalog: Dict[str, Dict[str, Any]]
) -> List[Dict[str, Any]]:
    """
    Ingest a single document file, extract text per page/slide, and attach metadata.

    Args:
        file_path: Path to target document file.
        metadata_catalog: Preloaded metadata catalog lookup dict.

    Returns:
        List of dictionary page/slide document records.
    """
    filename = file_path.name
    ext = file_path.suffix.lower()

    # Match metadata by filename or stem
    meta = metadata_catalog.get(
        filename.lower(),
        metadata_catalog.get(
            file_path.stem.lower(),
            {
                "document_id": file_path.stem,
                "filename": filename,
                "unit": 1,
                "topic": "General OS",
                "subtopic": "General",
                "source_type": "faculty_material",
                "source": "Local File Discovery",
                "license_or_access": "Institutional Internal Use",
                "description": f"Auto-discovered document {filename}",
            },
        ),
    )

    extracted_pages = extract_document(file_path)
    document_records: List[Dict[str, Any]] = []

    for item_num, text in extracted_pages:
        page_key = "slide" if ext == ".pptx" else "page"
        rec = {
            "document_id": meta["document_id"],
            "filename": filename,
            "source_path": str(file_path),
            page_key: item_num,
            "unit": meta["unit"],
            "topic": meta["topic"],
            "subtopic": meta["subtopic"],
            "source_type": meta["source_type"],
            "source": meta["source"],
            "license_or_access": meta["license_or_access"],
            "description": meta["description"],
            "raw_text": text,
        }
        document_records.append(rec)

    return document_records


def ingest_all_documents(
    raw_dir: Path = RAW_DATA_DIR,
    metadata_csv_path: Path = METADATA_CSV_PATH,
    taxonomy_json_path: Path = TAXONOMY_JSON_PATH,
) -> List[Dict[str, Any]]:
    """
    Main entry point: Discovers, ingests, and attaches metadata for all raw documents.

    Returns:
        Flattened list of document page/slide records.
    """
    catalog = load_and_validate_metadata(metadata_csv_path, taxonomy_json_path)
    discovered_paths = discover_documents(raw_dir)

    all_page_records: List[Dict[str, Any]] = []
    for p in discovered_paths:
        try:
            records = ingest_document(p, catalog)
            all_page_records.extend(records)
            logger.info(f"Ingested {p.name}: {len(records)} pages/slides extracted.")
        except Exception as e:
            logger.error(f"Failed to ingest document {p}: {e}")

    return all_page_records


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    print("Ingesting all raw documents...")
    docs = ingest_all_documents()
    print(f"Total document page/slide records extracted: {len(docs)}")
